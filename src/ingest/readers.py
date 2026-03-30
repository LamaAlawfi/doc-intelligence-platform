import os
from typing import Tuple, Dict, Any
from pypdf import PdfReader
from docx import Document

def read_file(file_path: str) -> Tuple[str, Dict[str, Any]]:
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    meta = {"source": os.path.basename(file_path), "path": file_path}
    
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            meta["pages"] = len(reader.pages)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        elif ext == ".docx":
            doc = Document(file_path)
            meta["pages"] = len(doc.paragraphs) // 20 + 1 # Rough estimate
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            meta["pages"] = len(text) // 3000 + 1
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    except Exception as e:
        raise Exception(f"Failed to read '{os.path.basename(file_path)}': {str(e)}")
    
    return text.strip(), meta
