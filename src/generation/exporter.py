import os
import io
import json
import arabic_reshaper
from bidi.algorithm import get_display
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fpdf import FPDF

# Configure Reshaper for Arabic
reshaper_config = {
    'delete_harakat': True,
    'shift_harakat_position': False,
    'pass_controls': False
}
reshaper = arabic_reshaper.ArabicReshaper(configuration=reshaper_config)

def process_arabic_text(text: str) -> str:
    """Reshapes and applies BiDi algorithm for proper Arabic rendering."""
    if not text or not text.strip():
        return text
    reshaped_text = reshaper.reshape(text)
    return get_display(reshaped_text)



class GovPDF(FPDF):
    def __init__(self, title, org_name, logo_path=None, theme_rgb=(0, 163, 108)):
        super().__init__()
        self.doc_title = title
        self.org_name = org_name
        self.logo_path = logo_path
        self.theme_rgb = theme_rgb
        
        font_dir = os.path.join(os.path.dirname(__file__), "..", "assets", "fonts")
        self.amiri_reg = os.path.join(font_dir, "Amiri-Regular.ttf")
        self.amiri_bold = os.path.join(font_dir, "Amiri-Bold.ttf")
        
        try:
            self.add_font('Amiri', '', self.amiri_reg, uni=True)
            self.add_font('Amiri', 'B', self.amiri_bold, uni=True)
            self.font_family = 'Amiri'
        except:
            self.font_family = 'Helvetica'
            
        self.set_margins(20, 20, 20)
        self.set_auto_page_break(auto=True, margin=20)

    def header(self):
        if self.page_no() == 1:
            return # No header on cover page
            
        is_curr = "Teacher Guide" in self.doc_title or "Student File" in self.doc_title or "دليل المعلم" in self.doc_title or "ملف الطالب" in self.doc_title
        if not is_curr:
            self.set_fill_color(*self.theme_rgb)
            self.rect(0, 0, 210, 5, 'F')
        else:
            self.set_fill_color(200, 200, 200)
            self.rect(10, 5, 190, 0.5, 'F')
        
        self.set_font(self.font_family, '', 10)
        self.set_text_color(100, 100, 100)
        
        # Logo or Org Name
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                self.image(self.logo_path, x=180, y=8, w=20)
            except:
                pass
        
        self.cell(0, 10, process_arabic_text(self.org_name), border=0, ln=0, align='L')
        self.ln(15)

    def footer(self):
        self.set_y(-15)
        self.set_font(self.font_family, '', 9)
        self.set_text_color(150, 150, 150)
        page_text = f"Page {self.page_no()}" if self.lang == "English" else f"الصفحة {self.page_no()}"
        self.cell(0, 10, process_arabic_text(page_text), 0, 0, 'C')

    def make_cover(self, metadata):
        self.add_page()
        
        is_curriculum = "Teacher Guide" in self.doc_title or "Student File" in self.doc_title or "دليل المعلم" in self.doc_title or "ملف الطالب" in self.doc_title
        
        if is_curriculum:
            self.set_text_color(0, 0, 0)
            self.set_y(40)
            
            curr_logo = self.logo_path if self.logo_path and os.path.exists(self.logo_path) else "storage/moe_logo.png"
            if os.path.exists(curr_logo):
                try:
                    self.image(curr_logo, x=85, y=30, w=40)
                    self.set_y(80)
                except:
                    self.set_y(60)
            else:
                self.set_y(60)
                
            self.set_font(self.font_family, 'B', 24)
            org_str = metadata.get('organization', 'وزارة التعليم') if metadata else 'وزارة التعليم'
            self.multi_cell(0, 15, process_arabic_text(org_str), align='C')
            self.ln(10)
            
            self.set_font(self.font_family, 'B', 28)
            self.multi_cell(0, 20, process_arabic_text(self.doc_title), align='C')
            
            self.ln(20)
            self.set_font(self.font_family, 'B', 18)
            self.set_text_color(150, 50, 50) 
            self.multi_cell(0, 10, process_arabic_text("Educational Document / وثيقة تعليمية"), align='C')
            
            self.ln(10)
            self.set_font(self.font_family, '', 16)
            self.set_text_color(0, 0, 0)
            self.multi_cell(0, 10, process_arabic_text(datetime.now().strftime('%Y/%m/%d')), align='C')
            
            return # Skip tables for curriculum
            
        # Normal Cover
        self.set_fill_color(*self.theme_rgb)
        self.rect(0, 0, 210, 40, 'F')
        
        self.set_y(50)
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                self.image(self.logo_path, x=85, y=50, w=40)
                self.set_y(100)
            except:
                pass
                
        self.set_font(self.font_family, 'B', 24)
        self.set_text_color(*self.theme_rgb)
        self.multi_cell(0, 15, process_arabic_text(self.doc_title), align='C')
        
        self.ln(20)
        self.set_text_color(0, 0, 0)
        
        # Info Table
        self.set_font(self.font_family, 'B', 14)
        info_title = "Document Information:" if self.lang == "English" else "معلومات الوثيقة:"
        self.cell(0, 10, process_arabic_text(info_title), ln=1, align='R' if self.lang == "Arabic" else 'L')
        
        self.set_font(self.font_family, '', 12)
        if self.lang == "English":
            info_data = [
                ("Author", metadata.get('organization', self.org_name)),
                ("Department", metadata.get('department', 'Government Entity')),
                ("Date", datetime.now().strftime('%Y/%m/%d')),
                ("Document Type", metadata.get('doc_type', 'Strategic Document')),
                ("Classification", metadata.get('classification', 'INTERNAL'))
            ]
        else:
            info_data = [
                ("معد الوثيقة", metadata.get('organization', self.org_name)),
                ("الجهة", metadata.get('department', 'الجهة الحكومية')),
                ("تاريخ الإعداد", datetime.now().strftime('%Y/%m/%d')),
                ("نوع الوثيقة", metadata.get('doc_type', 'وثيقة استراتيجية')),
                ("التصنيف", metadata.get('classification', 'INTERNAL'))
            ]
        
        self.set_fill_color(*self.theme_rgb)
        self.set_text_color(255, 255, 255)
        for k, v in info_data:
            if self.lang == "Arabic":
                self.cell(100, 10, process_arabic_text(v), border=1, align='R')
                self.cell(70, 10, process_arabic_text(k), border=1, align='R', fill=True)
            else:
                self.cell(70, 10, process_arabic_text(k), border=1, align='L', fill=True)
                self.cell(100, 10, process_arabic_text(v), border=1, align='L')
            self.ln()
            
        self.ln(20)
        
        # Version Table
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, 'B', 14)
        v_title = "Document Revision History:" if self.lang == "English" else "تاريخ إصدارات الوثيقة:"
        self.cell(0, 10, process_arabic_text(v_title), ln=1, align='R' if self.lang == "Arabic" else 'L')
        
        self.set_fill_color(*self.theme_rgb)
        self.set_text_color(255, 255, 255)
        self.set_font(self.font_family, 'B', 12)
        
        if self.lang == "Arabic":
            self.cell(70, 10, process_arabic_text("التغييرات بالوثيقة"), border=1, align='C', fill=True)
            self.cell(50, 10, process_arabic_text("معد الوثيقة"), border=1, align='C', fill=True)
            self.cell(30, 10, process_arabic_text("تاريخ الإصدار"), border=1, align='C', fill=True)
            self.cell(20, 10, process_arabic_text("الإصدار"), border=1, align='C', fill=True)
        else:
            self.cell(20, 10, process_arabic_text("Ver."), border=1, align='C', fill=True)
            self.cell(30, 10, process_arabic_text("Date"), border=1, align='C', fill=True)
            self.cell(50, 10, process_arabic_text("Author"), border=1, align='C', fill=True)
            self.cell(70, 10, process_arabic_text("Changes"), border=1, align='C', fill=True)
        self.ln()
        
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, '', 12)
        if self.lang == "Arabic":
            self.cell(70, 12, process_arabic_text("الإصدار الأول"), border=1, align='C')
            self.cell(50, 12, process_arabic_text(metadata.get('organization', self.org_name)), border=1, align='C')
            self.cell(30, 12, process_arabic_text(datetime.now().strftime('%Y/%m/%d')), border=1, align='C')
            self.cell(20, 12, process_arabic_text("0.1"), border=1, align='C')
        else:
            self.cell(20, 12, process_arabic_text("0.1"), border=1, align='C')
            self.cell(30, 12, process_arabic_text(datetime.now().strftime('%Y/%m/%d')), border=1, align='C')
            self.cell(50, 12, process_arabic_text(metadata.get('organization', self.org_name)), border=1, align='C')
            self.cell(70, 12, process_arabic_text("Initial Release"), border=1, align='C')
        self.ln()

    def make_curriculum_overview(self):
        self.add_page()
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, 'B', 24)
        self.cell(0, 15, process_arabic_text("أهداف المنهج"), ln=1, align='C')
        self.ln(10)
        
        safe_w = self.w - self.l_margin - self.r_margin

        # Helper to draw sections
        def draw_section(title, points):
            self.set_font(self.font_family, 'B', 18)
            self.set_text_color(*self.theme_rgb)
            self.cell(0, 12, process_arabic_text(title), ln=1, align='R')
            self.set_font(self.font_family, '', 14)
            self.set_text_color(0, 0, 0)
            for pt in points:
                # Naive RTL simulation for bullet points
                pt_text = "• " + pt
                words = pt_text.split()
                current_line = words[0]
                for word in words[1:]:
                    test_line = current_line + " " + word
                    if self.get_string_width(process_arabic_text(test_line)) > safe_w:
                        self.cell(0, 8, process_arabic_text(current_line), border=0, ln=1, align='R')
                        current_line = word
                    else:
                        current_line = test_line
                if current_line:
                    self.cell(0, 8, process_arabic_text(current_line), border=0, ln=1, align='R')
            self.ln(8)

        draw_section("الأهداف العامة", [
            "تزويد الطلاب بالمعرفة الأساسية والمتقدمة في الذكاء الاصطناعي.",
            "تمكين الطلاب من التفاعل الفعال والآمن مع نماذج اللغات الكبيرة.",
            "تعزيز التفكير النقدي وحل المشكلات باستخدام التقنيات الحديثة."
        ])
        
        draw_section("المهارات المستهدفة", [
            "صياغة الأوامر الدقيقة (Prompting) للحصول على أفضل النتائج.",
            "تحليل وتقييم جودة مخرجات الذكاء الاصطناعي التوليدي.",
            "التفكير المنطقي وبناء تسلسلات فعالة في الحوار مع الروبوتات."
        ])
        
        draw_section("نواتج التعلم", [
            "يميز الطالب بين الأوامر الضعيفة والقوية.",
            "يطبق الطالب استراتيجيات تحسين الأوامر لإنجاز مهام دراسية واقعية.",
            "يبني الطالب مشاريع مدرسية متكاملة مساعدة بالذكاء الاصطناعي."
        ])

    def make_importance_overview(self):
        self.add_page()
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, 'B', 24)
        self.cell(0, 15, process_arabic_text("أهمية هندسة الأوامر في العصر الحديث"), ln=1, align='C')
        self.ln(10)
        
        safe_w = self.w - self.l_margin - self.r_margin

        def draw_section(title, points):
            self.set_font(self.font_family, 'B', 18)
            self.set_text_color(*self.theme_rgb)
            self.cell(0, 12, process_arabic_text(title), ln=1, align='R')
            self.set_font(self.font_family, '', 14)
            self.set_text_color(0, 0, 0)
            for pt in points:
                pt_text = "• " + pt
                words = pt_text.split()
                current_line = words[0]
                for word in words[1:]:
                    test_line = current_line + " " + word
                    if self.get_string_width(process_arabic_text(test_line)) > safe_w:
                        self.cell(0, 8, process_arabic_text(current_line), border=0, ln=1, align='R')
                        current_line = word
                    else:
                        current_line = test_line
                if current_line:
                    self.cell(0, 8, process_arabic_text(current_line), border=0, ln=1, align='R')
            self.ln(6)

        draw_section("الذكاء الاصطناعي في التعليم", [
            "يُستخدم لتحسين وتخصيص تجربة التعلم عبر تحفيز التفكير النقدي والذاتي.",
            "مساعدة الطلاب في تلخيص النصوص، مراجعة الدروس، وابتكار أفكار للمشاريع الدراسية."
        ])
        
        draw_section("الذكاء الاصطناعي في سوق العمل", [
            "أصبحت مهارات صياغة الأوامر مطلبًا رئيسيًا وتنافسيًا في الوظائف الحديثة.",
            "يساهم بشكل مباشر في تسريع أداء المهام كصياغة رسائل البريد الإلكتروني والتقارير."
        ])
        
        draw_section("التفاعل الصحيح والآمن مع الآلة", [
            "القدرة على وضع أسئلة دقيقة تُنتج إجابات واضحة وخالية من التخيلات (Hallucinations).",
            "تحليل وتقييم جودة مخرجات الذكاء الاصطناعي لضمان تفادي الأخطاء وعدم الاعتماد الأعمى."
        ])
        
        draw_section("الملاءمة للتوجهات المستقبلية", [
            "تأهيل جيل قادر على المنافسة في العصر الرقمي المتسارع.",
            "الانتقال من الاستهلاك التقني البسيط إلى مرحلة القيادة والتوجيه الفعّال لخوارزميات الذكاء الاصطناعي."
        ])

    def make_definitions_table(self, definitions_text):
        if not definitions_text: return
        self.add_page()
        self.set_font(self.font_family, 'B', 18)
        self.set_text_color(*self.theme_rgb)
        title = "Definitions and Acronyms" if self.lang == "English" else "التعريفات والاختصارات"
        self.cell(0, 15, process_arabic_text(title), ln=1, align='C')
        self.ln(10)
        
        self.set_fill_color(*self.theme_rgb)
        self.set_text_color(255, 255, 255)
        self.set_font(self.font_family, 'B', 12)
        
        if self.lang == "Arabic":
            self.cell(100, 10, process_arabic_text("الوصف (Definition)"), border=1, align='C', fill=True)
            self.cell(50, 10, process_arabic_text("الاختصار (Acronyms)"), border=1, align='C', fill=True)
            self.cell(20, 10, process_arabic_text("#"), border=1, align='C', fill=True)
        else:
            self.cell(20, 10, process_arabic_text("#"), border=1, align='C', fill=True)
            self.cell(50, 10, process_arabic_text("Acronym"), border=1, align='C', fill=True)
            self.cell(100, 10, process_arabic_text("Definition"), border=1, align='C', fill=True)
        self.ln()
        
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, '', 12)
        
        # Robust parsing: handle ':', '-', or just list format
        import re
        lines = [l.strip() for l in definitions_text.split('\n') if l.strip()]
        for i, line in enumerate(lines):
            # Try to split by ':' or '-'
            parts = re.split(r'[:\-]', line, maxsplit=1)
            acronym = parts[0].strip() if len(parts) > 0 else (f"مصطلح {i+1}" if self.lang == "Arabic" else f"Term {i+1}")
            desc = parts[1].strip() if len(parts) > 1 else acronym
            
            # Use multi_cell for description to allow wrapping
            curr_y = self.get_y()
            # Calculate height for multi_cell first
            # We use a dummy multi_cell or calculate based on string width
            # For simplicity, we just render and then draw the other cells with the same height
            if self.lang == "Arabic":
                # Draw desc first to get height
                self.set_xy(10, curr_y)
                self.multi_cell(100, 10, process_arabic_text(desc), border=1, align='R')
                new_y = self.get_y()
                h = new_y - curr_y
                # Draw others with height h
                self.set_xy(110, curr_y)
                self.cell(50, h, process_arabic_text(acronym), border=1, align='C')
                self.set_xy(160, curr_y)
                self.cell(20, h, process_arabic_text(str(i+1)), border=1, align='C')
                self.set_y(new_y)
            else:
                self.set_xy(160, curr_y) # Temporary move
                # Actually simpler order for English
                desc_text = process_arabic_text(desc)
                # Render desc and get height
                self.set_xy(80, curr_y)
                self.multi_cell(100, 10, desc_text, border=1, align='L')
                new_y = self.get_y()
                h = new_y - curr_y
                # Render others
                self.set_xy(10, curr_y)
                self.cell(20, h, process_arabic_text(str(i+1)), border=1, align='C')
                self.set_xy(30, curr_y)
                self.cell(50, h, process_arabic_text(acronym), border=1, align='C')
                self.set_y(new_y)

    def make_review_approval_tables(self, reviewers_text, approvers_text):
        if not reviewers_text and not approvers_text: return
        self.add_page()
        
        def draw_table(title, raw_text):
            self.set_font(self.font_family, 'B', 16)
            self.set_text_color(*self.theme_rgb)
            self.cell(0, 10, process_arabic_text(title), ln=1, align='R' if self.lang == "Arabic" else 'L')
            self.ln(5)
            self.set_fill_color(*self.theme_rgb)
            self.set_text_color(255, 255, 255)
            self.set_font(self.font_family, 'B', 11)
            
            headers = ["Name", "Job Title", "Signature", "Date"] if self.lang == "English" else ["الاسم", "المسمى الوظيفي", "التوقيع", "التاريخ"]
            widths = [50, 50, 40, 30]
            if self.lang == "Arabic":
                headers.reverse()
                widths.reverse()
            
            for i, h in enumerate(headers):
                self.cell(widths[i], 10, process_arabic_text(h), border=1, align='C', fill=True)
            self.ln()
            
            self.set_text_color(0, 0, 0)
            self.set_font(self.font_family, '', 11)
            
            import re
            # Split by newline or semicolon to isolate individual people
            person_blocks = re.split(r'[\n;]', raw_text)
            
            # Fallback: If only one block but it contains multiple colons/hyphens, 
            # the user might have written "Name: Title : Name: Title"
            if len(person_blocks) == 1:
                block = person_blocks[0]
                if block.count(':') >= 2 or block.count('-') >= 2:
                    # Try to separate by whatever delimiter has more than 2 hits
                    delim = ':' if block.count(':') >= 2 else '-'
                    sub_parts = block.split(delim)
                    # pair them up: [Name, Title, Name, Title...]
                    person_blocks = []
                    for i in range(0, len(sub_parts), 2):
                        if i + 1 < len(sub_parts):
                            person_blocks.append(f"{sub_parts[i].strip()} {delim} {sub_parts[i+1].strip()}")
                        else:
                            person_blocks.append(sub_parts[i].strip())

            for block in person_blocks:
                block = block.strip()
                if not block: continue
                
                # Split Name from Title using '-' or ':'
                parts = re.split(r'[:\-]', block, maxsplit=1)
                name = parts[0].strip() if len(parts) > 0 else "N/A"
                title = parts[1].strip() if len(parts) > 1 else ("مسؤول" if self.lang == "Arabic" else "Official")
                
                if self.lang == "Arabic":
                    self.cell(30, 12, "", border=1)
                    self.cell(40, 12, "", border=1)
                    self.cell(50, 12, process_arabic_text(title), border=1, align='R')
                    self.cell(50, 12, process_arabic_text(name), border=1, align='R')
                else:
                    self.cell(50, 12, process_arabic_text(name), border=1, align='L')
                    self.cell(50, 12, process_arabic_text(title), border=1, align='L')
                    self.cell(40, 12, "", border=1)
                    self.cell(30, 12, "", border=1)
                self.ln()
            self.ln(10)

        if reviewers_text:
            draw_table("Review / المراجعة" if self.lang == "English" else "المراجعة", reviewers_text)
        if approvers_text:
            draw_table("Approval / الاعتماد" if self.lang == "English" else "الاعتماد", approvers_text)

    def make_toc(self, sections):
        self.add_page()
        self.set_font(self.font_family, 'B', 20)
        self.set_text_color(*self.theme_rgb)
        title = "Table of Contents" if self.lang == "English" else "جدول المحتويات"
        self.cell(0, 15, process_arabic_text(title), ln=1, align='C')
        self.ln(10)
        self.set_text_color(0, 0, 0)
        self.set_font(self.font_family, '', 14)
        
        for k, title in sections.items():
            tit = process_arabic_text(title)
            if self.lang == "English":
                self.cell(0, 10, tit + " " + "." * 60, ln=1, align='L')
            else:
                # For Arabic: visually we want [Dots][Title] so that reading RTL it's [Title][Dots]
                dotted = ("." * 60) + " " + tit
                self.cell(0, 10, dotted, ln=1, align='R')

class Exporter:
    @staticmethod
    def _write_arabic_paragraph(pdf, text: str, safe_w: float, line_height: int = 8, align: str = 'R'):
        """
        Custom word-wrapper for Arabic text. FPDF multi_cell breaks on visually-reversed RTL text.
        This function logically breaks the raw string into lines that fit the width,
        THEN reverses and renders each line individually.
        """
        # Clean raw markdown
        text = text.replace('**', '').replace('__', '')
        if text.startswith('- ') or text.startswith('* '):
            text = "• " + text[2:]
            
        words = text.split()
        if not words:
            return
            
        current_line = words[0]
        for word in words[1:]:
            test_line = current_line + " " + word
            # Measure using the reversed (visual) string because that matches font glyphs
            width = pdf.get_string_width(process_arabic_text(test_line))
            
            if width > safe_w:
                pdf.cell(0, line_height, process_arabic_text(current_line), border=0, ln=1, align=align)
                current_line = word
            else:
                current_line = test_line
                
        if current_line:
            pdf.cell(0, line_height, process_arabic_text(current_line), border=0, ln=1, align=align)

    @staticmethod
    def to_docx(title: str, content_md: str, metadata: dict = None) -> str:
        doc = Document()
        styles = doc.styles
        style = styles['Normal']
        font = style.font
        font.name = 'Arial'
        
        if metadata:
            t = doc.add_heading(title, 0)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(f"Organization: {metadata.get('organization', 'Official Government System')}\n")
            p.add_run(f"Department: {metadata.get('department', 'General Administration')}\n")
            p.add_run(f"Date: {datetime.now().strftime('%Y-%m-%d')}\n")
            doc.add_page_break()

        lines = content_md.split('\n')
        for line in lines:
            line = line.strip()
            if not line: continue
            
            p = None
            if line.startswith('###'): 
                p = doc.add_heading(line.replace('###', '').strip(), 3)
            elif line.startswith('##'): 
                p = doc.add_heading(line.replace('##', '').strip(), 2)
            elif line.startswith('#'): 
                p = doc.add_heading(line.replace('#', '').strip(), 1)
            else: 
                p = doc.add_paragraph(line)
            
            if p: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        os.makedirs("storage/exports", exist_ok=True)
        filename = f"document_{datetime.now().strftime('%Y%p%d_%H%M%S')}.docx"
        file_path = os.path.join("storage/exports", filename)
        doc.save(file_path)
        return file_path

    @staticmethod
    def to_pdf(title: str, content_md: str, metadata: dict = None, template_struct: dict = None, logo_path: str = None) -> str:
        org_name = metadata.get('organization', 'المؤسسة الحكومية') if metadata else 'المؤسسة الحكومية'
        
        # Override Theme for Curriculum
        is_curriculum = "Teacher Guide" in title or "Student File" in title or "دليل المعلم" in title or "ملف الطالب" in title
        if is_curriculum:
            theme_rgb = (50, 50, 50) # Neutral Dark Gray
        else:
            theme_str = str(metadata.get('color_theme', 'green')).lower() if metadata else 'green'
            theme_rgb = (0, 163, 108) # Default Green
            if 'red' in theme_str or 'أحمر' in theme_str:
                theme_rgb = (180, 20, 30) # Professional Red
            elif 'blue' in theme_str or 'أزرق' in theme_str:
                theme_rgb = (20, 80, 180) # Professional Blue
            elif 'gray' in theme_str or 'رمادي' in theme_str:
                theme_rgb = (90, 90, 90)  # Professional Gray
            
        pdf = GovPDF(title, org_name, logo_path, theme_rgb=theme_rgb)
        pdf.lang = metadata.get('language', 'Arabic') if metadata else 'Arabic'
        
        try:
            # Metadata could be None, avoid crushing
            meta = metadata if metadata else {}
            pdf.make_cover(meta)
            
            # Extract sections for TOC from template_struct
            if template_struct:
                sections = {k: v.get('title', k) for k, v in template_struct.items()}
                pdf.make_toc(sections)
                
            # New Sections: Definitions and Acronyms
            if meta.get('definitions_acronyms'):
                pdf.make_definitions_table(meta['definitions_acronyms'])
                
            # Content
            pdf.add_page()
            pdf.set_font(pdf.font_family, "", 12)
            pdf.set_text_color(0, 0, 0)
            safe_w = pdf.w - pdf.l_margin - pdf.r_margin

            # Language alignment
            main_align = 'L' if pdf.lang == "English" else 'R'

            lines = content_md.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    pdf.ln(5)
                    continue
                
                # Check for bilingual separators
                if "--- ARABIC" in line or "--- ENGLISH" in line:
                    lang_tag = "Arabic" if "ARABIC" in line else "English"
                    pdf.set_font(pdf.font_family, "B", 10)
                    pdf.set_text_color(120, 120, 120)
                    pdf.cell(0, 10, process_arabic_text(line), ln=1, align='C')
                    pdf.set_font(pdf.font_family, "", 12)
                    pdf.set_text_color(0, 0, 0)
                    continue

                if line.startswith('#'):
                    level = 1
                    if line.startswith('###'): level = 3
                    elif line.startswith('##'): level = 2
                    
                    title_text = line.lstrip('#').strip().lower()
                    # Force page break for Intro and Conclusion
                    if any(x in title_text for x in ["مقدمة", "introduction", "الخاتمة", "conclusion", "خاتمة"]):
                        pdf.add_page()
                    
                    font_size = 18 if level == 1 else (16 if level == 2 else 14)
                    pdf.set_font(pdf.font_family, "B", font_size)
                    
                    if level == 2:
                       pdf.set_text_color(*pdf.theme_rgb) # Theme-colored headers for H2
                    else:
                       pdf.set_text_color(0, 0, 0)
                       
                    pdf.ln(8)
                    Exporter._write_arabic_paragraph(pdf, line.lstrip('#').strip(), safe_w, 10, align=main_align)
                    pdf.set_font(pdf.font_family, "", 12)
                    pdf.set_text_color(0, 0, 0)
                else:
                    try:
                        # Auto-detect alignment for bilingual lines?
                        # For now, use main_align, but bilingual intro/concl might need special care
                        curr_align = main_align
                        if pdf.lang == "Arabic" and any(c.isalpha() for c in line) and not any('\u0600' <= c <= '\u06FF' for c in line):
                            curr_align = 'L' # fallback for English text in Arabic doc
                        Exporter._write_arabic_paragraph(pdf, line, safe_w, 8, align=curr_align)
                    except Exception as e:
                        pdf.write(5, f"[Error: {e}]")
                        pdf.ln(7)

            # Review and Approval tables
            pdf.make_review_approval_tables(meta.get('reviewers'), meta.get('approvers'))

        except Exception as e:
            print("PDF generation caught an internal exception:", e)

        os.makedirs("storage/exports", exist_ok=True)
        filename = f"document_{datetime.now().strftime('%Y%p%d_%H%M%S')}.pdf"
        file_path = os.path.join("storage/exports", filename)
        pdf.output(file_path)
        return file_path
